"""URL matching against links/REBRANDLY.xlsx for CUTTLY and DOB.CO SEO sheets."""

from __future__ import annotations

from dataclasses import dataclass
import os
from pathlib import Path
import sys
from typing import Optional
import openpyxl

# ANSI color codes - Bold High-Intensity Bright Green
ANSI_BOLD_GREEN = "\033[1;92m"
ANSI_RESET = "\033[0m"

# Enable ANSI escape sequences on Windows console host if needed
if sys.platform == "win32":
    try:
        os.system("")
    except Exception:
        pass


@dataclass(frozen=True)
class RebrandlyMatch:
    """Represents a matching URL found in REBRANDLY.xlsx."""

    brand: str
    source_type: str  # 'CUTTLY' or 'DOB.CO'
    sheet_name: str   # 'CUTTLY' or 'DOB.CO SEO'
    sheet_url: str    # Exact URL in Excel
    matched_url: str  # The URL that was matched

    def format_message(self, colorize: bool = True) -> str:
        """Format the identification alert message."""
        msg = f"{self.brand} {self.source_type} URL identified: {self.matched_url}"
        if colorize:
            return f"{ANSI_BOLD_GREEN}{msg}{ANSI_RESET}"
        return msg


def normalize_lookup_key(url: str, strip_query: bool = True, lowercase_path: bool = False) -> str:
    """Normalize a URL to a canonical key for robust matching.

    1. Trims whitespace.
    2. Strips protocol (http://, https://, //).
    3. Strips URL fragments (#...).
    4. Optionally strips query parameters (?ref=...).
    5. Normalizes host to lowercase.
    6. Strips trailing slashes from path.
    7. Optionally lowercases path.
    """
    url = str(url).strip()
    if not url:
        return ""
    if "://" in url:
        _, rest = url.split("://", 1)
    elif url.startswith("//"):
        rest = url[2:]
    else:
        rest = url

    # Remove fragment
    if "#" in rest:
        rest = rest.split("#", 1)[0]

    # Optionally remove query string
    if strip_query and "?" in rest:
        rest = rest.split("?", 1)[0]

    if "/" in rest:
        host, path = rest.split("/", 1)
        path = path.rstrip("/")
        path = ("/" + path) if path else ""
    else:
        host = rest
        path = ""

    host = host.lower()
    if lowercase_path:
        path = path.lower()

    return f"{host}{path}"


class RebrandlyMatcher:
    """Loads and matches URLs against CUTTLY and DOB.CO SEO sheets in REBRANDLY.xlsx."""

    def __init__(self, excel_path: str | Path = "links/REBRANDLY.xlsx"):
        self.excel_path = Path(excel_path)
        self.exact_map: dict[str, dict] = {}
        self.lower_map: dict[str, dict] = {}
        self.loaded = False
        self._load()

    def _load(self) -> None:
        if not self.excel_path.is_file():
            print(f"[RebrandlyMatcher] Warning: Spreadsheet not found at {self.excel_path}", flush=True)
            return

        try:
            wb = openpyxl.load_workbook(self.excel_path, data_only=True, read_only=True)
            sheet_configs = [
                ("CUTTLY", "CUTTLY", ["WEB"], ["LINK CUTT.LY"]),
                ("DOB.CO SEO", "DOB.CO", ["BRAND"], ["LINK DOB.CO"]),
            ]

            for sheet_name, source_type, brand_headers, link_headers in sheet_configs:
                if sheet_name not in wb.sheetnames:
                    continue
                ws = wb[sheet_name]
                header = None
                brand_col, link_col = 0, 1

                for row in ws.iter_rows(values_only=True):
                    if not row or not any(row):
                        continue
                    if header is None:
                        header = [str(c).strip().upper() if c is not None else "" for c in row]
                        for idx, h in enumerate(header):
                            if any(bh in h for bh in brand_headers):
                                brand_col = idx
                            if any(lh in h for lh in link_headers):
                                link_col = idx
                        continue

                    if len(row) <= max(brand_col, link_col):
                        continue

                    brand_val = row[brand_col]
                    link_val = row[link_col]
                    if not brand_val or not link_val:
                        continue

                    brand_str = str(brand_val).strip()
                    link_str = str(link_val).strip()

                    # Skip header repeat if any
                    if link_str.upper() in {"LINK CUTT.LY", "LINK DOB.CO", "LINK"}:
                        continue

                    info = {
                        "brand": brand_str,
                        "source_type": source_type,
                        "sheet_name": sheet_name,
                        "sheet_url": link_str,
                    }

                    # Exact path key (with query stripped)
                    k_exact = normalize_lookup_key(link_str, strip_query=True, lowercase_path=False)
                    if k_exact and k_exact not in self.exact_map:
                        self.exact_map[k_exact] = info

                    # Lowercase path key
                    k_lower = normalize_lookup_key(link_str, strip_query=True, lowercase_path=True)
                    if k_lower and k_lower not in self.lower_map:
                        self.lower_map[k_lower] = info

            wb.close()
            self.loaded = True
            print(f"[RebrandlyMatcher] Loaded {len(self.exact_map)} unique link keys from {self.excel_path.name}.", flush=True)
        except Exception as exc:
            print(f"[RebrandlyMatcher] Error loading spreadsheet: {exc}", flush=True)

    def check_url(self, url: str) -> Optional[RebrandlyMatch]:
        """Check if a given URL matches any link in REBRANDLY.xlsx."""
        if not self.loaded or not url:
            return None

        # 1. Try exact path key (query stripped)
        k_exact = normalize_lookup_key(url, strip_query=True, lowercase_path=False)
        info = self.exact_map.get(k_exact)

        # 2. Try with query if query was present
        if not info and "?" in url:
            k_with_q = normalize_lookup_key(url, strip_query=False, lowercase_path=False)
            info = self.exact_map.get(k_with_q)

        # 3. Fallback to lowercase path
        if not info:
            k_lower = normalize_lookup_key(url, strip_query=True, lowercase_path=True)
            info = self.lower_map.get(k_lower)

        if info:
            return RebrandlyMatch(
                brand=info["brand"],
                source_type=info["source_type"],
                sheet_name=info["sheet_name"],
                sheet_url=info["sheet_url"],
                matched_url=url,
            )
        return None

    def format_match_message(self, match: RebrandlyMatch, colorize: bool = True) -> str:
        """Format the match message using bold bright green."""
        return match.format_message(colorize=colorize)
